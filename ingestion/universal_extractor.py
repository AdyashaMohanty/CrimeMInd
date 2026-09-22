"""
CrimeMind Universal Evidence Extractor

Automatically detects supported file formats and extracts their contents
into a common Python structure.

Usage:
    python -m ingestion.universal_extractor <file_or_directory>
"""

from pathlib import Path
import sys
import json
import csv
import xml.etree.ElementTree as ET
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
import zipfile


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".docx",
    ".csv",
    ".xlsx",
    ".json",
    ".xml",
    ".eml",
    ".vcf",
    ".ics",
    ".srt",
    ".log",
    ".html",
    ".kml",
    ".geojson",
    ".jpg",
    ".jpeg",
    ".png",
    ".wav",
    ".mp4",
}


# ============================================================
# TXT / LOG / SRT
# ============================================================

def extract_text(path):
    return {
        "type": "text",
        "content": path.read_text(
            encoding="utf-8",
            errors="replace"
        )
    }


# ============================================================
# DOCX
# ============================================================

def extract_docx(path):
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError(
            "python-docx is required. Install with: pip install python-docx"
        )

    # Some files in the synthetic world may have a .docx extension
    # but contain plain text. Handle that safely.
    try:
        document = Document(path)

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        tables = []

        for table in document.tables:
            rows = []

            for row in table.rows:
                rows.append([
                    cell.text for cell in row.cells
                ])

            tables.append(rows)

        return {
            "type": "docx",
            "paragraphs": paragraphs,
            "tables": tables,
        }

    except (zipfile.BadZipFile, Exception) as exc:
        # Fallback for malformed/non-standard synthetic DOCX files
        try:
            text = path.read_text(
                encoding="utf-8",
                errors="replace"
            )

            return {
                "type": "docx_text_fallback",
                "content": text,
                "warning": f"DOCX parsing failed: {exc}",
            }

        except Exception:
            raise


# ============================================================
# CSV
# ============================================================

def extract_csv(path):
    with open(
        path,
        "r",
        encoding="utf-8",
        errors="replace",
        newline=""
    ) as f:

        reader = csv.DictReader(f)

        rows = list(reader)

        return {
            "type": "csv",
            "columns": reader.fieldnames or [],
            "rows": rows,
            "row_count": len(rows),
        }


# ============================================================
# XLSX
# ============================================================

def extract_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError(
            "openpyxl is required. Install with: pip install openpyxl"
        )

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True
    )

    sheets = {}

    for worksheet in workbook.worksheets:

        rows = []

        for row in worksheet.iter_rows(values_only=True):
            rows.append(list(row))

        sheets[worksheet.title] = rows

    workbook.close()

    return {
        "type": "xlsx",
        "sheets": sheets,
    }


# ============================================================
# JSON
# ============================================================

def extract_json(path):
    with open(
        path,
        "r",
        encoding="utf-8",
        errors="replace"
    ) as f:

        data = json.load(f)

    return {
        "type": "json",
        "data": data,
    }


# ============================================================
# XML
# ============================================================

def xml_element_to_dict(element):
    children = list(element)

    if not children:
        return (element.text or "").strip()

    result = {}

    for child in children:
        value = xml_element_to_dict(child)

        if child.tag in result:

            if not isinstance(result[child.tag], list):
                result[child.tag] = [
                    result[child.tag]
                ]

            result[child.tag].append(value)

        else:
            result[child.tag] = value

    return result


def extract_xml(path):
    tree = ET.parse(path)
    root = tree.getroot()

    return {
        "type": "xml",
        "root": root.tag,
        "data": {
            root.tag: xml_element_to_dict(root)
        },
    }


# ============================================================
# EML
# ============================================================

def get_email_body(message):

    if message.is_multipart():

        parts = []

        for part in message.walk():

            if part.get_content_type() == "text/plain":

                try:
                    parts.append(part.get_content())
                except Exception:
                    pass

        return "\n".join(parts).strip()

    try:
        return message.get_content()
    except Exception:
        return ""


def extract_eml(path):

    with open(path, "rb") as f:
        message = BytesParser(
            policy=policy.default
        ).parse(f)

    return {
        "type": "eml",
        "from": message.get("From"),
        "to": message.get("To"),
        "cc": message.get("Cc"),
        "subject": message.get("Subject"),
        "date": message.get("Date"),
        "body": get_email_body(message),
    }


# ============================================================
# VCF
# ============================================================

def extract_vcf(path):

    text = path.read_text(
        encoding="utf-8",
        errors="replace"
    )

    contact = {}

    for line in text.splitlines():

        if ":" not in line:
            continue

        key, value = line.split(":", 1)

        key = key.split(";", 1)[0].upper()

        if key == "FN":
            contact["name"] = value

        elif key == "TEL":
            contact["phone"] = value

        elif key == "EMAIL":
            contact["email"] = value

        elif key == "ORG":
            contact["organization"] = value

    return {
        "type": "vcf",
        "contact": contact,
    }


# ============================================================
# ICS
# ============================================================

def extract_ics(path):

    text = path.read_text(
        encoding="utf-8",
        errors="replace"
    )

    event = {}

    mapping = {
        "SUMMARY": "summary",
        "DESCRIPTION": "description",
        "DTSTART": "start",
        "DTEND": "end",
        "LOCATION": "location",
        "UID": "uid",
        "ORGANIZER": "organizer",
    }

    for line in text.splitlines():

        if ":" not in line:
            continue

        key, value = line.split(":", 1)

        key = key.split(";", 1)[0].upper()

        if key in mapping:
            event[mapping[key]] = value

    return {
        "type": "ics",
        "event": event,
    }


# ============================================================
# HTML
# ============================================================

class HTMLTextParser(HTMLParser):

    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_data(self, data):
        if data.strip():
            self.parts.append(data.strip())


def extract_html(path):

    html = path.read_text(
        encoding="utf-8",
        errors="replace"
    )

    parser = HTMLTextParser()
    parser.feed(html)

    return {
        "type": "html",
        "content": "\n".join(parser.parts),
    }


# ============================================================
# KML
# ============================================================

def extract_kml(path):

    tree = ET.parse(path)
    root = tree.getroot()

    places = []

    for element in root.iter():

        if element.tag.endswith("Placemark"):

            place = {}

            for child in element.iter():

                tag = child.tag.split("}")[-1]

                if tag in {
                    "name",
                    "description",
                    "coordinates",
                }:

                    place[tag] = (
                        child.text or ""
                    ).strip()

            if place:
                places.append(place)

    return {
        "type": "kml",
        "places": places,
    }


# ============================================================
# GEOJSON
# ============================================================

def extract_geojson(path):

    with open(
        path,
        "r",
        encoding="utf-8",
        errors="replace"
    ) as f:

        data = json.load(f)

    return {
        "type": "geojson",
        "data": data,
    }


# ============================================================
# IMAGE
# ============================================================

def extract_image(path):

    result = {
        "type": "image",
        "filename": path.name,
        "extension": path.suffix.lower(),
        "size_bytes": path.stat().st_size,
    }

    try:
        from PIL import Image

        image = Image.open(path)

        result["width"] = image.width
        result["height"] = image.height
        result["format"] = image.format

    except ImportError:

        result["warning"] = (
            "Pillow not installed; image metadata limited."
        )

    except Exception as exc:

        result["warning"] = str(exc)

    return result


# ============================================================
# WAV
# ============================================================

def extract_wav(path):

    import wave

    with wave.open(str(path), "rb") as audio:

        return {
            "type": "wav",
            "channels": audio.getnchannels(),
            "sample_width": audio.getsampwidth(),
            "frame_rate": audio.getframerate(),
            "frame_count": audio.getnframes(),
            "duration_seconds": (
                audio.getnframes()
                / audio.getframerate()
                if audio.getframerate()
                else 0
            ),
        }


# ============================================================
# MP4
# ============================================================

def extract_mp4(path):

    return {
        "type": "mp4",
        "filename": path.name,
        "size_bytes": path.stat().st_size,
        "status": "CCTV media detected; frame extraction handled separately.",
    }


# ============================================================
# DISPATCHER
# ============================================================

def extract_file(file_path):

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"File does not exist: {path}"
        )

    if not path.is_file():
        raise ValueError(
            f"Not a file: {path}"
        )

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    extractors = {

        ".txt": extract_text,
        ".log": extract_text,
        ".srt": extract_text,

        ".docx": extract_docx,

        ".csv": extract_csv,

        ".xlsx": extract_xlsx,

        ".json": extract_json,

        ".xml": extract_xml,

        ".eml": extract_eml,

        ".vcf": extract_vcf,

        ".ics": extract_ics,

        ".html": extract_html,

        ".kml": extract_kml,

        ".geojson": extract_geojson,

        ".jpg": extract_image,
        ".jpeg": extract_image,
        ".png": extract_image,

        ".wav": extract_wav,

        ".mp4": extract_mp4,
    }

    return extractors[extension](path)


# ============================================================
# DIRECTORY PROCESSING
# ============================================================

def scan_directory(directory):

    directory = Path(directory)

    if not directory.exists():
        raise FileNotFoundError(
            f"Directory does not exist: {directory}"
        )

    files = [
        file
        for file in directory.rglob("*")
        if file.is_file()
        and file.suffix.lower() in SUPPORTED_EXTENSIONS
    ]

    return sorted(files)


def process_directory(directory):

    files = scan_directory(directory)

    print(f"Found {len(files)} supported files.\n")

    results = []

    success = 0
    failed = 0

    for index, file in enumerate(files, start=1):

        relative = file.relative_to(directory)

        print(
            f"[{index}/{len(files)}] "
            f"{relative}"
        )

        try:

            extracted = extract_file(file)

            results.append({
                "source_file": str(relative),
                "extension": file.suffix.lower(),
                "success": True,
                "data": extracted,
            })

            success += 1

            print("    ✓ extracted")

        except Exception as exc:

            results.append({
                "source_file": str(relative),
                "extension": file.suffix.lower(),
                "success": False,
                "error": str(exc),
            })

            failed += 1

            print(
                f"    ✗ failed: {exc}"
            )

    print("\n========================================")
    print("EXTRACTION SUMMARY")
    print("========================================")
    print(f"Total files : {len(files)}")
    print(f"Successful  : {success}")
    print(f"Failed      : {failed}")
    print("========================================")

    return results


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":

    if len(sys.argv) != 2:

        print(
            "Usage:\n"
            "  python -m ingestion.universal_extractor <file>\n"
            "  python -m ingestion.universal_extractor <directory>"
        )

        raise SystemExit(1)

    target = Path(sys.argv[1])

    if target.is_dir():

        results = process_directory(target)

        # Save extracted results for the normalization stage.
        output_file = Path(
            "data/extracted_evidence.json"
        )

        output_file.parent.mkdir(
            parents=True,
            exist_ok=True
        )

        with open(
            output_file,
            "w",
            encoding="utf-8"
        ) as f:

            json.dump(
                results,
                f,
                indent=2,
                ensure_ascii=False,
                default=str
            )

        print(
            f"\nSaved extraction output to:\n"
            f"{output_file}"
        )

    else:

        result = extract_file(target)

        print("\n--- EXTRACTED DATA ---")

        print(
            json.dumps(
                result,
                indent=2,
                ensure_ascii=False,
                default=str
            )
        )