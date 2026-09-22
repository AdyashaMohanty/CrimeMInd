from pathlib import Path
from openpyxl import load_workbook


def extract_xlsx(file_path: str):
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File does not exist: {path}")

    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"Expected .xlsx file, got {path.suffix}")

    workbook = load_workbook(path, data_only=True)

    sheets = {}

    for worksheet in workbook.worksheets:
        rows = list(
            worksheet.iter_rows(
                values_only=True
            )
        )

        sheets[worksheet.title] = rows

    workbook.close()

    return sheets


if __name__ == "__main__":
    import sys

    if len(sys.argv) != 2:
        print(
            "Usage: python -m ingestion.extractors.xlsx_extractor <file>"
        )
        raise SystemExit(1)

    sheets = extract_xlsx(sys.argv[1])

    print(f"Extracted {len(sheets)} worksheet(s).")

    for sheet_name, rows in sheets.items():
        print(f"\n===== SHEET: {sheet_name} =====")
        print(f"Rows: {len(rows)}")

        for row in rows[:5]:
            print(row)